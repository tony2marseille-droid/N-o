import csv
import io
from pathlib import Path


class ExtractionError(Exception):
    pass


MAX_CHARS = 400_000


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    if suffix == ".csv":
        return _extract_csv(path)
    return _extract_plain(path)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("pypdf manquant") from exc
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(f"[Page {i + 1}]\n{text}")
    result = "\n\n".join(pages)
    return result[:MAX_CHARS]


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("python-docx manquant") from exc
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)[:MAX_CHARS]


def _extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ExtractionError("openpyxl manquant") from exc
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for sheet in wb.worksheets:
        out.append(f"[Feuille: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)[:MAX_CHARS]


def _extract_csv(path: Path) -> str:
    raw = path.read_text("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    lines = [" | ".join(row) for row in reader]
    return "\n".join(lines)[:MAX_CHARS]


def _extract_plain(path: Path) -> str:
    return path.read_text("utf-8", errors="replace")[:MAX_CHARS]


def chunk_text(text: str, target: int = 900, overlap: int = 120) -> list[str]:
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for para in paragraphs:
        if len(para) > target * 2:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(para):
                piece = para[start : start + target]
                chunks.append(piece)
                start += target - overlap
            continue
        if len(current) + len(para) + 1 > target:
            chunks.append(current)
            tail = current[-overlap:] if overlap < len(current) else current
            current = f"{tail}\n{para}"
        else:
            current = f"{current}\n{para}" if current else para
    if current.strip():
        chunks.append(current)
    return [c.strip() for c in chunks if c.strip()]
